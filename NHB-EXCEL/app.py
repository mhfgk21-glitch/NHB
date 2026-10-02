import streamlit as st
import openpyxl
import shutil
import os
import threading
from datetime import datetime
import colorsys

# Page config
st.set_page_config(
    page_title=" نسمة هوى بغداد - Excel Template Filler",
    page_icon="🧾",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- App Constants ---
APP_NAME = "نسمة هوى بغداد"
APP_VERSION = "1.0.0"
APP_AUTHOR = "محمد هادي"
APP_COPYRIGHT = "© 2025 محمد هادي - جميع الحقوق محفوظة"

NAMES_LIST = [
    "الشارقة", "القمة", "كوول", "المحطة", "الرائد",
    "المحيط", "بريق النجوم", "الياسمين", "اليمامة", 
    "الجود", "الكوثر", "الريان", "كول"
]

COLORS = {
    "primary": "#2E86AB",
    "secondary": "#A23B72",
    "success": "#18A558",
    "warning": "#F39C12",
    "danger": "#E74C3C",
    "info": "#3498DB",
    "dark": "#2C3E50",
    "light_bg": "#ECF0F1",
    "section_bg": "white",
    "accent1": "#9B59B6",
    "accent2": "#1ABC9C",
    "accent3": "#E67E22",
    "accent4": "#E74C3C",
    "entry_bg": "#F8F9FA",
}

# --- License Manager (simplified for web) ---
class LicenseManager:
    def __init__(self):
        self.license_data = None
    
    def get_hardware_id(self):
        return "HWID-DEMO-001"
    
    def save_license(self, license_key, data):
        if len(license_key) == 8:
            self.license_data = data
            return True
        return False
    
    def is_licensed(self):
        if self.license_data:
            return True, "مرخص"
        return False, "غير مفعل"
    
    def get_license_info(self):
        if self.license_data:
            return {
                "الحالة": "مفعل",
                "العميل": self.license_data.get("customer", ""),
                "الإصدار": self.license_data.get("version", "")
            }
        return None

# --- Column Mapping ---
COLUMN_MAPPING = {
    "رقم الوصل": ["رقم الوصل"],
    "رقم الهاتف": ["رقم الهاتف", "هاتف المستلم", "هاتف"],
    "العنوان": ["العنوان"],
    "المبلغ": ["المبلغ", "المبلغ المطلوب", "السعر"]
}

# --- Template Config Manager ---
class TemplateConfigManager:
    def __init__(self, config_file="template_config.json"):
        self.config_file = config_file
        self.configs = self.load_configs()
    
    def load_configs(self):
        if os.path.exists(self.config_file):
            try:
                with open(self.config_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception as e:
                print(f"خطأ في تحميل التكوينات: {e}")
                return {}
        return {}
    
    def save_config(self, company_name, header_row, column_mapping):
        self.configs[company_name] = {
            "header_row": header_row,
            "columns": column_mapping
        }
        try:
            with open(self.config_file, 'w', encoding='utf-8') as f:
                json.dump(self.configs, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            print(f"خطأ فيSaving التكوين: {e}")
            return False
    
    def get_config(self, company_name):
        return self.configs.get(company_name)
    
    def detect_template(self, ws, filename=None, max_search=20):
        potential_matches = []
        for i in range(1, max_search + 1):
            row = ws[i]
            headers = [
                " ".join(
                    str(cell.value)
                    .strip()
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .split()
                )
                for cell in row if cell.value is not None
            ]
            if not headers:
                continue
            for company_name, config in self.configs.items():
                if config["header_row"] != i:
                    continue
                required_columns = set(config["columns"].values())
                found_columns = set(headers)
                if required_columns.issubset(found_columns):
                    potential_matches.append((company_name, config))
        if not potential_matches:
            return None, None
        if len(potential_matches) == 1:
            return potential_matches[0]
        if filename:
            filename_clean = os.path.splitext(os.path.basename(filename))[0]
            for name, config in potential_matches:
                if name in filename_clean or filename_clean in name:
                    return name, config
        return potential_matches[0]
    
    def get_all_companies(self):
        return list(self.configs.keys())

# --- Excel Processor ---
class ExcelProcessor:
    def __init__(self, template_path, column_mapping):
        self.template_path = template_path
        self.column_mapping = column_mapping

    def find_header_row(self, ws, max_search=20):
        for i in range(1, max_search + 1):
            row = ws[i]
            headers = [
                " ".join(
                    str(cell.value)
                    .strip()
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .split()
                )
                for cell in row if cell.value is not None
            ]
            found_mapping = {}
            all_found = True
            for canonical, aliases in self.column_mapping.items():
                match = next((h for h in headers if h in aliases), None)
                if match:
                    found_mapping[canonical] = match
                else:
                    all_found = False
                    break
            if all_found:
                return i, headers, found_mapping
        return None, None, None

    def process_file(self, data_file, out_folder, selected_name, update_callback, stop_flag_ref, config_manager=None):
        try:
            data_wb = openpyxl.load_workbook(data_file)
            data_ws = data_wb.active
        except Exception as e:
            raise Exception(f"خطأ في تحميل ملف البيانات: {e}")

        detected_config = None
        if config_manager:
            _, detected_config = config_manager.detect_template(data_ws, filename=data_file)
        
        if detected_config:
            required_columns = set(detected_config["columns"].values())
            header_row_num = detected_config["header_row"]
            row = data_ws[header_row_num]
            headers = [
                " ".join(
                    str(cell.value)
                    .strip()
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .split()
                )
                for cell in row if cell.value is not None
            ]
            found_columns = set(headers)
            if not required_columns.issubset(found_columns):
                raise ValueError(f"لم يتم العثور على الأعمدة المطلوبة في الصف {header_row_num}: {required_columns - found_columns}")
            found_map = {}
            for canonical, actual_name in detected_config["columns"].items():
                if actual_name in headers:
                    found_map[canonical] = actual_name
            header_row = header_row_num
        else:
            header_row, headers, found_map = self.find_header_row(data_ws)
            if header_row is None:
                required_names = [f"{k} ({'/'.join(v)})" for k, v in self.column_mapping.items()]
                raise ValueError(f"لم يتم العثور على صف يحتوي الأعمدة المطلوبة: {', '.join(required_names)}")

        output_filename = f"{selected_name}.xlsx"
        output_path = os.path.join(out_folder, output_filename)
        
        try:
            shutil.copy(self.template_path, output_path)
            out_wb = openpyxl.load_workbook(output_path)
            out_ws = out_wb.active
        except Exception as e:
            raise Exception(f"خطأ في إعداد ملف الإخراج: {e}")
        
        start_row = 2
        count = 0
        total = data_ws.max_row - header_row
        
        if total <= 0:
            raise ValueError("لا توجد سجلات بيانات لمعالجتها في ملف الإدخال.")
            
        for i, row in enumerate(data_ws.iter_rows(min_row=header_row + 1, values_only=True)):
            if stop_flag_ref():
                break

            record = dict(zip(headers, row))
            
            num = str(record.get(found_map["رقم الوصل"], "") or "").strip()
            if not num:
                continue
            
            phone = str(record.get(found_map["رقم الهاتف"], "") or "").strip()
            address = str(record.get(found_map["العنوان"], "") or "").strip()
            raw_amount = record.get(found_map["المبلغ"])
            amount = str(raw_amount).strip() if raw_amount is not None else ""
            
            product_type = ""
            quantity = ""
            notes = ""
            
            for header_name in headers:
                if "نوع البضاعة" in header_name or "نوع" in header_name:
                    product_type = str(record.get(header_name, "") or "").strip()
                elif "العدد" in header_name or "عدد" in header_name:
                    quantity = str(record.get(header_name, "") or "").strip()
                elif "ملاحظات" in header_name or "الملاحظات" in header_name:
                    notes = str(record.get(header_name, "") or "").strip()
            
            region = address.replace("محافظة ذي قار", "").replace("ذي قار", "").replace("-", "").strip()
            governorate = "ذي قار"
            if selected_name == "الجود":
                governorate = "الناصرية"
            
            if selected_name == " كوول":
                num = num.replace("ANT_", "")
            
            if selected_name == "الرائد":
                governorate = "ذي قار"
            
            if not region:
                region = "الناصرية"
            
            current_row = start_row + count
            out_ws.cell(row=current_row, column=1).value = num
            out_ws.cell(row=current_row, column=3).value = phone
            out_ws.cell(row=current_row, column=5).value = governorate
            out_ws.cell(row=current_row, column=6).value = region
            out_ws.cell(row=current_row, column=7).value = amount
            
            if product_type:
                out_ws.cell(row=current_row, column=8).value = product_type
            if quantity:
                out_ws.cell(row=current_row, column=9).value = quantity
            if notes:
                out_ws.cell(row=current_row, column=10).value = notes

            count += 1
            if update_callback:
                update_callback(count, total, selected_name)
        
        if not stop_flag_ref():
            try:
                out_wb.save(output_path)
            except Exception as e:
                raise Exception(f"خطأ في حفظ الملف الناتج: تأكد من إغلاق الملف: {e}")

        return count, output_path


# --- Streamlit App ---
class ExcelTemplateAppStreamlit:
    def __init__(self):
        self.config_manager = TemplateConfigManager()
        self.processor = ExcelProcessor("PM.xlsx", COLUMN_MAPPING)
        self.license_manager = LicenseManager()
        self.stop_flag = False
        self.reset_state()
    
    def reset_state(self):
        self.data_path = st.session_state.get('data_path', "")
        self.output_folder = st.session_state.get('output_folder', os.path.expanduser("~/Desktop"))
        self.selected_name = st.session_state.get('selected_name', NAMES_LIST[0])
        self.progress_val = st.session_state.get('progress_val', 0)
        self.status_text = st.session_state.get('status_text', "جاهز للبدء...")
    
    def darken_color(self, color, factor=0.2):
        try:
            color = color.lstrip('#')
            rgb = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
            h, l, s = colorsys.rgb_to_hls(rgb[0]/255, rgb[1]/255, rgb[2]/255)
            l = max(0, l - factor)
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            return '#{:02x}{:02x}{:02x}'.format(int(r*255), int(g*255), int(b*255))
        except:
            return color
    
    def run(self):
        st.title(" 🧾 ")
        st.subheader("نظام تعبئة القوالب - Excel")
        
        # License check
        if st.session_state.get('license_checked', False) is False:
            is_valid, message = self.license_manager.is_licensed()
            if not is_valid:
                st.warning(f"⚠️ {message}")
                license_key = st.text_input("مفتاح الترخيص", type="password", key="license_key")
                if st.button("تفعيل") and license_key:
                    if self.license_manager.save_license(license_key, {
                        "hardware_id": self.license_manager.get_hardware_id(),
                        "customer": "Web User",
                        "created": datetime.now().isoformat(),
                        "expires": "PERMANENT",
                        "version": APP_VERSION
                    }):
                        st.session_state.license_checked = True
                        st.success("✅ تم التفعيل بنجاح")
                        st.rerun()
                    else:
                        st.error("❌ مفتاح غير صحيح (يجب أن يكون 8 رموز)")
                return
        
        # Main interface
        col1, col2 = st.columns([2, 1])
        
        with col1:
            st.markdown("### 📁 تحميل ملف البيانات")
            uploaded_file = st.file_uploader("اختر ملف Excel", type=["xlsx", "xls"], key="file_uploader")
            
            if uploaded_file:
                # Save uploaded file temporarily
                tmp_path = os.path.join(".", uploaded_file.name)
                with open(tmp_path, "wb") as f:
                    f.write(uploaded_file.getbuffer())
                
                # Detect template
                try:
                    wb = openpyxl.load_workbook(tmp_path)
                    ws = wb.active
                    company_name, config = self.config_manager.detect_template(ws, filename=tmp_path)
                    
                    if company_name:
                        st.success(f"✅ تم اكتشاف قالب: {company_name}")
                        if company_name in NAMES_LIST:
                            st.session_state.selected_name = company_name
                    
                    # Show preview
                    st.markdown("### المعاينة (أول 5 صفوف)")
                    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
                        with st.expander(f"صف {i+1}", expanded=False):
                            st.write([str(v) if v is not None else "" for v in row])
                except Exception as e:
                    st.error(f"خطأ في الملف: {str(e)}")
                finally:
                    if os.path.exists(tmp_path):
                        os.remove(tmp_path)
        
        with col2:
            st.markdown("### 👤 اختيار الاسم")
            selected_name = st.selectbox("اسم الشركة", NAMES_LIST, index=NAMES_LIST.index(st.session_state.selected_name) if st.session_state.selected_name in NAMES_LIST else 0, key="name_select")
            st.session_state.selected_name = selected_name
        
        # Output folder
        st.markdown("### ⚙️ مجلد الحفظ")
        output_folder = st.text_input("مسار الحفظ", value=st.session_state.output_folder or os.path.expanduser("~/Desktop"), key="output_folder_input")
        st.session_state.output_folder = output_folder
        
        # Start button
        col_start, col_stop = st.columns([1, 1])
        
        with col_start:
            if st.button("🚀 ابدأ التنفيذ", type="primary", use_container_width=True):
                if not st.session_state.get('file_uploaded', False):
                    st.warning("يرجى اختيار ملف بيانات أولاً")
                elif not output_folder:
                    st.warning("يرجى تحديد مجلد الحفظ")
                else:
                    self.run_process(output_folder)
        
        with col_stop:
            if st.button("⏹️ إيقاف", use_container_width=True):
                self.stop_flag = True
                st.stop_flag = True
        
        # Progress
        st.progress(st.session_state.get('progress_val', 0) / max(1, st.session_state.get('progress_total', 1)))
        st.caption(st.session_state.get('status_text', "جاهز للبدء..."))
    
    def run_process(self, output_folder):
        self.stop_flag = False
        data_file = st.session_state.get('file_uploaded_path', '')
        selected_name = st.session_state.selected_name
        
        if not data_file or not os.path.exists(data_file):
            st.error("ملف البيانات غير موجود")
            return
        
        # Reset progress
        st.session_state.progress_val = 0
        st.session_state.progress_total = 1
        st.session_state.status_text = "جاري التحميل..."
        st.rerun()
        
        def update_callback(count, total, name):
            st.session_state.progress_val = count
            st.session_state.progress_total = total
            st.session_state.status_text = f"⏳ جاري معالجة {count}/{total} 👤 {name}"
            st.rerun()
        
        def stop_flag_ref():
            return self.stop_flag
        
        with st.spinner("جاري المعالجة..."):
            try:
                count, output_path = self.processor.process_file(
                    data_file=data_file,
                    out_folder=output_folder,
                    selected_name=selected_name,
                    update_callback=update_callback,
                    stop_flag_ref=stop_flag_ref,
                    config_manager=self.config_manager
                )
                
                st.session_state.progress_val = count
                st.session_state.progress_total = count if count > 0 else 1
                st.session_state.status_text = f"✅ تم الانتهاء! {count} سجل 👤 {selected_name}"
                st.success(f"🎉 تم إنشاء الملف!\n📁 المسار: {output_path}\n📊 عدد السجلات: {count}\n👤 الاسم: {selected_name}")
                
                # Offer download
                with open(output_path, "rb") as f:
                    st.download_button(
                        label="📥 تحميل الملف الناتج",
                        data=f.read(),
                        file_name=f"{selected_name}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.session_state.status_text = f"❌ خطأ: {str(e)}"
                st.error(f"حدث خطأ: {str(e)}")
            finally:
                st.rerun()


if __name__ == "__main__":
    app = ExcelTemplateAppStreamlit()
    app.run()