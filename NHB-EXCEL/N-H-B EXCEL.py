# Copyright (c) 2025 محمد هادي - جميع الحقوق محفوظة
# Unauthorized copying, distribution, or modification is prohibited

import os
import shutil
import openpyxl
import customtkinter as ctk
from tkinterdnd2 import DND_FILES, TkinterDnD
from tkinter import filedialog, messagebox
import tkinter as tk
import threading
from datetime import datetime
import colorsys
import json

# نظام الترخيص والحماية
try:
    from license_manager import LicenseManager
    LICENSE_ENABLED = True  # مفعّل للنسخة النهائية
except ImportError:
    LICENSE_ENABLED = False
    print("تحذير: نظام الترخيص غير متوفر")

class AppConfig:
    """تجميع الثوابت وإعدادات الألوان."""
    
    # معلومات البرنامج
    APP_NAME = "نسمة هوى بغداد"
    APP_VERSION = "1.0.0"
    APP_BUILD_DATE = "2025-12-04"
    APP_AUTHOR = "محمد هادي"
    APP_COPYRIGHT = "© 2025   محمد هادي- جميع الحقوق محفوظة"
    
    DEFAULT_THEME = "blue"
    APPEARANCE_MODE = "Light"
    WINDOW_SIZE = "800x550"
    TEMPLATE_PATH = "PM.xlsx"
    FONT_FAMILY = "Arial" 

    COLORS = {
        "primary": "#2E86AB",       # أزرق غامق مميز (للهيدر والخلفيات)
        "secondary": "#A23B72",     # بنفسجي (لإظهار بعض العناصر)
        "success": "#18A558",       # أخضر (لبدء العملية)
        "warning": "#F39C12",       # برتقالي (لإعدادات الإخراج)
        "danger": "#E74C3C",        # أحمر (لإيقاف العملية)
        "info": "#3498DB",          # أزرق فاتح (لتحميل الملف)
        "dark": "#2C3E50",          # أزرق غامق جداً (للنصوص الداكنة)
        "light_bg": "#ECF0F1",      # فاتح جداً (الخلفية الرئيسية)
        "section_bg": "white",      # خلفية الأقسام
        "accent1": "#9B59B6",       # أرجواني (شريط التقدم)
        "accent2": "#1ABC9C",       # فيروزي (اختيار الاسم)
        "accent3": "#E67E22",       # برتقالي داكن (لوحة التحكم)
        "accent4": "#E74C3C",       # أحمر مرجاني (إعادة تعيين)
        "entry_bg": "#F8F9FA",      # خلفية حقول الإدخال
    }
    
    # قائمة الأسماء
    NAMES_LIST = [
        "الشارقة", "القمة", "كوول", "المحطة", "الرائد",
        "المحيط", "بريق النجوم", "الياسمين", "اليمامة", 
        "الجود", "الكوثر", "الريان", "كول"
    ]

    # إعدادات القوائم المنسدلة (Dropdown Style)
    DROPDOWN_STYLE = {
        "height": 42,
        "corner_radius": 12,
        "border_width": 2,
        "font": ("Arial", 13, "bold"),
        "dropdown_font": ("Arial", 12),
        "fg_color": "#F8F9FA",
        "dropdown_fg_color": "#FFFFFF",
        "dropdown_text_color": "#2C3E50",
        "dropdown_hover_color": "#E3F2FD",
        "text_color": "#2C3E50",
    }

    # خريطة الأعمدة المطلوبة للبحث في ملف الإدخال
    COLUMN_MAPPING = {
        "رقم الوصل": ["رقم الوصل"],
        "رقم الهاتف": ["رقم الهاتف", "هاتف المستلم", "هاتف"],
        "العنوان": ["العنوان"],
        "المبلغ": ["المبلغ", "المبلغ المطلوب", "السعر"]
    }

# --- 1.5 إدارة تكوينات القوالب ---
class TemplateConfigManager:
    """إدارة حفظ واسترجاع تكوينات القوالب المختلفة."""
    
    def __init__(self, config_file="template_config.json"):
        self.config_file = config_file
        self.configs = self.load_configs()
    
    def load_configs(self):
        """تحميل التكوينات من ملف JSON."""
        if os.path.exists(self.config_file):
            try:
                with open(self.config_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception as e:
                print(f"خطأ في تحميل التكوينات: {e}")
                return {}
        return {}
    
    def save_config(self, company_name, header_row, column_mapping):
        """حفظ تكوين جديد لشركة."""
        self.configs[company_name] = {
            "header_row": header_row,
            "columns": column_mapping
        }
        
        try:
            with open(self.config_file, 'w', encoding='utf-8') as f:
                json.dump(self.configs, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            print(f"خطأ في حفظ التكوين: {e}")
            return False
    
    def get_config(self, company_name):
        """الحصول على تكوين شركة معينة."""
        return self.configs.get(company_name)
    
    def detect_template(self, ws, filename=None, max_search=20):
        """محاولة التعرف على القالب من بنية الملف واسمه."""
        # جمع معلومات عن بنية الملف
        potential_matches = []

        for i in range(1, max_search + 1):
            row = ws[i]
            headers = [
                " ".join(
                    str(cell.value)
                    .strip()
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .split()
                )
                for cell in row if cell.value is not None
            ]
            
            if not headers:
                continue
            
            # مقارنة مع كل تكوين محفوظ
            for company_name, config in self.configs.items():
                if config["header_row"] != i:
                    continue
                
                # التحقق من تطابق الأعمدة
                required_columns = set(config["columns"].values())
                found_columns = set(headers)
                
                # إذا كانت جميع الأعمدة المطلوبة موجودة
                if required_columns.issubset(found_columns):
                    potential_matches.append((company_name, config))
        
        if not potential_matches:
            return None, None
        
        # إذا وجدنا تطابق واحد فقط
        if len(potential_matches) == 1:
            return potential_matches[0]
        
        # إذا وجدنا أكثر من تطابق، نستخدم اسم الملف للترجيح
        if filename:
            filename_clean = os.path.splitext(os.path.basename(filename))[0]
            for name, config in potential_matches:
                # تنظيف الاسم للمقارنة (إزالة المسافات والهمزات قد يكون مفيداً)
                if name in filename_clean or filename_clean in name:
                    return name, config
        
        return potential_matches[0]
    
    def get_all_companies(self):
        """الحصول على قائمة بجميع الشركات المحفوظة."""
        return list(self.configs.keys())

class ExcelProcessor:
    """فئة مخصصة لمعالجة ملفات Excel والبحث عن الأعمدة."""

    def __init__(self, template_path, column_mapping):
        self.template_path = template_path
        self.column_mapping = column_mapping

    def find_header_row(self, ws, max_search=20):
        """يبحث عن صف الهيدر الذي يحتوي على جميع الأعمدة المطلوبة."""
        for i in range(1, max_search + 1):
            row = ws[i]
            # تنظيف قيم الهيدر
            headers = [
                " ".join(
                    str(cell.value)
                    .strip()
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .split()
                )
                for cell in row if cell.value is not None
            ]
            
            # محاولة مطابقة جميع الأعمدة المطلوبة
            found_mapping = {}
            all_found = True
            for canonical, aliases in self.column_mapping.items():
                match = next((h for h in headers if h in aliases), None)
                if match:
                    # نستخدم الاسم الذي تم العثور عليه فعلاً
                    found_mapping[canonical] = match
                else:
                    all_found = False
                    break
            
            if all_found:
                # إرجاع رقم الصف، وقائمة الرؤوس، وخريطة المطابقة
                return i, headers, found_mapping
        
        return None, None, None

    def process_file(self, data_file, out_folder, selected_name, update_callback, stop_flag_ref, config_manager=None):
        """يقوم بتحميل ملف البيانات ومعالجته وتعبئة القالب."""
        
        # 1. تحميل ملفات Excel
        try:
            data_wb = openpyxl.load_workbook(data_file)
            data_ws = data_wb.active
        except Exception as e:
            raise Exception(f"خطأ في تحميل ملف البيانات: {e}")

        # 2. اكتشاف القالب للحصول على التكوين الصحيح
        detected_config = None
        if config_manager:
            _, detected_config = config_manager.detect_template(data_ws, filename=data_file)
        
        # 3. البحث عن صف الهيدر باستخدام التكوين المكتشف أو الافتراضي
        if detected_config:
            # استخدام الأعمدة المحددة من التكوين المكتشف
            required_columns = set(detected_config["columns"].values())
            header_row_num = detected_config["header_row"]
            
            # التحقق من أن الصف المحدد يحتوي على الأعمدة المطلوبة
            row = data_ws[header_row_num]
            headers = [
                " ".join(
                    str(cell.value)
                    .strip()
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .split()
                )
                for cell in row if cell.value is not None
            ]
            
            found_columns = set(headers)
            if not required_columns.issubset(found_columns):
                raise ValueError(f"لم يتم العثور على الأعمدة المطلوبة في الصف {header_row_num}: {required_columns - found_columns}")
            
            # إنشاء خريطة المطابقة
            found_map = {}
            for canonical, actual_name in detected_config["columns"].items():
                if actual_name in headers:
                    found_map[canonical] = actual_name
            
            header_row = header_row_num
        else:
            # استخدام المنطق القديم (generic mapping)
            header_row, headers, found_map = self.find_header_row(data_ws)
            if header_row is None:
                required_names = [f"{k} ({'/'.join(v)})" for k, v in self.column_mapping.items()]
                raise ValueError(f"لم يتم العثور على صف يحتوي الأعمدة المطلوبة: {', '.join(required_names)}")

        # 3. إنشاء اسم الملف الناتج بناءً على الاسم المختار
        output_filename = f"{selected_name}.xlsx"
        output_path = os.path.join(out_folder, output_filename)
        
        try:
            shutil.copy(self.template_path, output_path)
            out_wb = openpyxl.load_workbook(output_path)
            out_ws = out_wb.active
        except Exception as e:
            raise Exception(f"خطأ في إعداد ملف الإخراج: {e}")
            
        start_row = 2  # البدء من الصف الثاني مباشرة (بعد الهيدر)

        # 4. معالجة الصفوف
        count = 0
        total = data_ws.max_row - header_row
        
        if total <= 0:
            raise ValueError("لا توجد سجلات بيانات لمعالجتها في ملف الإدخال.")
            
        for i, row in enumerate(data_ws.iter_rows(min_row=header_row + 1, values_only=True)):
            if stop_flag_ref():
                break

            record = dict(zip(headers, row))
            
            # استخراج البيانات من صفوف الإدخال باستخدام الأعمدة التي تم العثور عليها
            num = str(record.get(found_map["رقم الوصل"], "") or "").strip()
            
            # تخطي الصفوف التي لا تحتوي على رقم وصل
            if not num:
                continue
            
            phone = str(record.get(found_map["رقم الهاتف"], "") or "").strip()
            address = str(record.get(found_map["العنوان"], "") or "").strip()
            # معالجة المبلغ (للحفاظ على الصفر)
            raw_amount = record.get(found_map["المبلغ"])
            amount = str(raw_amount).strip() if raw_amount is not None else ""
            
            # استخراج الحقول الاختيارية إذا كانت موجودة
            product_type = ""
            quantity = ""
            notes = ""
            
            # البحث عن الحقول الاختيارية في الملف المصدر
            for header_name in headers:
                if "نوع البضاعة" in header_name or "نوع" in header_name:
                    product_type = str(record.get(header_name, "") or "").strip()
                elif "العدد" in header_name or "عدد" in header_name:
                    quantity = str(record.get(header_name, "") or "").strip()
                elif "ملاحظات" in header_name or "الملاحظات" in header_name:
                    notes = str(record.get(header_name, "") or "").strip()
            
            # منطق معالجة العنوان (إزالة "ذي قار" و "-" لإيجاد المنطقة)
            region = address.replace("محافظة ذي قار", "").replace("ذي قار", "").replace("-", "").strip()

            # تحديد المحافظة
            governorate = "ذي قار"
            if selected_name == "الجود":
                governorate = "الناصرية"

            # معالجة خاصة للاسم "كوول" - حذف "ANT_" من رقم الوصل
            if selected_name == "كوول":
                num = num.replace("ANT_", "")

            # معالجة خاصة للاسم "الرائد"
            if selected_name == "الرائد":
                governorate = "ذي قار"
            
            # ملء المنطقة الفارغة بـ "الناصرية" لجميع الشركات
            if not region:
                region = "الناصرية"

            current_row = start_row + count  # استخدام count بدلاً من i
            # تعبئة خلايا الإخراج (الحقول الأساسية)
            out_ws.cell(row=current_row, column=1).value = num
            out_ws.cell(row=current_row, column=3).value = phone
            out_ws.cell(row=current_row, column=5).value = governorate
            out_ws.cell(row=current_row, column=6).value = region
            out_ws.cell(row=current_row, column=7).value = amount
            
            # تعبئة الحقول الاختيارية إذا كانت موجودة
            if product_type:
                out_ws.cell(row=current_row, column=8).value = product_type
            if quantity:
                out_ws.cell(row=current_row, column=9).value = quantity
            if notes:
                out_ws.cell(row=current_row, column=10).value = notes

            count += 1
            update_callback(count, total, selected_name) # تحديث التقدم وواجهة المستخدم

        # 5. حفظ الملف
        if not stop_flag_ref():
            try:
                out_wb.save(output_path)
            except Exception as e:
                raise Exception(f"خطأ في حفظ الملف الناتج: تأكد من إغلاق الملف: {e}")

        return count, output_path # إرجاع عدد السجلات ومسار الملف

# --- 2.5 نافذة تعلم القالب ---
class TemplateLearnWindow(ctk.CTkToplevel):
    """نافذة اضافة قالب جديد."""

    def __init__(self, parent, config_manager):
        super().__init__(parent)
        self.config_manager = config_manager
        
        self.title("📚 اضافة قالب جديد")
        self.geometry("800x650")  # تعديل الحجم ليكون مناسباً أكثر
        self.resizable(True, True) # السماح بتغيير الحجم
        
        # المتغيرات
        self.file_path = None
        self.wb = None
        self.ws = None
        self.headers = []
        self.header_row_num = 1
        
        self.company_name = ctk.StringVar()
        self.selected_receipt_col = ctk.StringVar(value="-- اختر العمود --")
        self.selected_phone_col = ctk.StringVar(value="-- اختر العمود --")
        self.selected_address_col = ctk.StringVar(value="-- اختر العمود --")
        self.selected_amount_col = ctk.StringVar(value="-- اختر العمود --")
        
        # الخطوط
        self.font_h1 = ctk.CTkFont(size=16, weight="bold")
        self.font_body = ctk.CTkFont(size=12, weight="bold")
        self.font_small = ctk.CTkFont(size=11)
        
        self.create_ui()
    
    def create_ui(self):
        """إنشاء واجهة المستخدم."""
        main_frame = ctk.CTkFrame(self, fg_color="#ECF0F1")
        main_frame.pack(fill="both", expand=True, padx=15, pady=15)
        
        # العنوان
        title_frame = ctk.CTkFrame(main_frame, fg_color="#2E86AB", corner_radius=10, height=50)
        title_frame.pack(fill="x", pady=(0, 10))
        title_frame.pack_propagate(False)
        
        ctk.CTkLabel(title_frame, text="📚 اضافة قالب جديد", 
                     font=self.font_h1, text_color="white").pack(pady=12)

        # تقسيم الشاشة إلى عمودين
        content_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        content_frame.pack(fill="both", expand=True)
        content_frame.columnconfigure(0, weight=1)
        content_frame.columnconfigure(1, weight=1)

        # === العمود الأيسر (تحميل ومعاينة) ===
        left_col = ctk.CTkFrame(content_frame, fg_color="transparent")
        left_col.grid(row=0, column=0, sticky="nsew", padx=(0, 5))

        # قسم تحميل الملف
        file_frame = ctk.CTkFrame(left_col, fg_color="white", corner_radius=10)
        file_frame.pack(fill="x", pady=(0, 10))
        
        ctk.CTkLabel(file_frame, text="1️⃣ اختر ملف نموذجي:", 
                     font=self.font_body, text_color="#2C3E50").pack(anchor="w", padx=10, pady=(10, 5))
        
        btn_frame = ctk.CTkFrame(file_frame, fg_color="transparent")
        btn_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkButton(btn_frame, text="📂 اختيار ملف Excel",
                      command=self.load_sample_file,
                      fg_color="#3498DB",
                      hover_color="#2980B9",
                      font=self.font_small).pack(side="left")
        
        self.file_label = ctk.CTkLabel(btn_frame, text="لم يتم اختيار ملف",
                                       font=self.font_small, text_color="#7F8C8D")
        self.file_label.pack(side="left", padx=10)
        
        # قسم معاينة البيانات
        preview_frame = ctk.CTkFrame(left_col, fg_color="white", corner_radius=10)
        preview_frame.pack(fill="both", expand=True)
        
        ctk.CTkLabel(preview_frame, text="2️⃣ معاينة البيانات (أول 5 صفوف):", 
                     font=self.font_body, text_color="#2C3E50").pack(anchor="w", padx=10, pady=(10, 5))
        
        # إطار قابل للتمرير للمعاينة
        self.preview_text = ctk.CTkTextbox(preview_frame, font=self.font_small)
        self.preview_text.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.preview_text.configure(state="disabled")

        # === العمود الأيمن (تحديد الأعمدة والحفظ) ===
        right_col = ctk.CTkFrame(content_frame, fg_color="transparent")
        right_col.grid(row=0, column=1, sticky="nsew", padx=(5, 0))

        # قسم تحديد الأعمدة
        mapping_frame = ctk.CTkFrame(right_col, fg_color="white", corner_radius=10)
        mapping_frame.pack(fill="x", pady=(0, 10))
        
        ctk.CTkLabel(mapping_frame, text="3️⃣ حدد الأعمدة المطلوبة:", 
                     font=self.font_body, text_color="#2C3E50").pack(anchor="w", padx=10, pady=(10, 5))
        
        cols_frame = ctk.CTkFrame(mapping_frame, fg_color="transparent")
        cols_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        # ترتيب عمودي للحقول في العمود الأيمن
        cols_frame.columnconfigure(0, weight=1)
        
        self.create_column_selector(cols_frame, "رقم الوصل:", self.selected_receipt_col, 0, 0)
        self.create_column_selector(cols_frame, "رقم الهاتف:", self.selected_phone_col, 1, 0)
        self.create_column_selector(cols_frame, "العنوان:", self.selected_address_col, 2, 0)
        self.create_column_selector(cols_frame, "المبلغ:", self.selected_amount_col, 3, 0)
        
        # قسم اسم الشركة والحفظ
        save_frame = ctk.CTkFrame(right_col, fg_color="white", corner_radius=10)
        save_frame.pack(fill="x", pady=(0, 0))
        
        ctk.CTkLabel(save_frame, text="4️⃣ اسم الشركة:", 
                     font=self.font_body, text_color="#2C3E50").pack(anchor="w", padx=10, pady=(10, 5))
        
        name_entry = ctk.CTkEntry(save_frame, textvariable=self.company_name,
                                  placeholder_text="أدخل اسم الشركة...",
                                  font=self.font_small, height=35)
        name_entry.pack(fill="x", padx=10, pady=(0, 10))
        
        # أزرار الحفظ والإلغاء
        btn_save_frame = ctk.CTkFrame(save_frame, fg_color="transparent")
        btn_save_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkButton(btn_save_frame, text="💾 حفظ التكوين",
                      command=self.save_configuration,
                      fg_color="#18A558",
                      hover_color="#16A34A",
                      font=self.font_body,
                      height=40).pack(side="left", padx=5, expand=True, fill="x")
        
        ctk.CTkButton(btn_save_frame, text="❌ إلغاء",
                      command=self.destroy,
                      fg_color="#E74C3C",
                      hover_color="#C0392B",
                      font=self.font_body,
                      height=40).pack(side="left", padx=5, expand=True, fill="x")
    
    def create_column_selector(self, parent, label_text, variable, row, col):
        """إنشاء محدد عمود."""
        frame = ctk.CTkFrame(parent, fg_color="transparent")
        frame.grid(row=row, column=col, padx=5, pady=5, sticky="ew")
        
        ctk.CTkLabel(frame, text=label_text, font=self.font_small,
                     text_color="#2C3E50").pack(anchor="w")
        
        dropdown = ctk.CTkComboBox(frame, variable=variable,
                                   values=["-- اختر العمود --"],
                                   state="readonly",
                                   **AppConfig.DROPDOWN_STYLE,
                                   border_color="#3498DB",  # Info color
                                   button_color="#3498DB",
                                   button_hover_color="#2980B9")
        dropdown.pack(fill="x", pady=(2, 0))
        
        return dropdown
    
    def load_sample_file(self):
        """تحميل ملف نموذجي."""
        file_path = filedialog.askopenfilename(
            title="اختر ملف Excel نموذجي",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        
        if not file_path:
            return
        
        try:
            self.file_path = file_path
            self.wb = openpyxl.load_workbook(file_path)
            self.ws = self.wb.active
            
            # البحث عن صف الهيدر
            self.find_and_display_headers()
            
            self.file_label.configure(text=f"✅ {os.path.basename(file_path)}")
            
        except Exception as e:
            messagebox.showerror("خطأ", f"فشل تحميل الملف:\n{str(e)}")
    
    def find_and_display_headers(self):
        """البحث عن صف الهيدر وعرض المعاينة."""
        # البحث عن صف الهيدر (أول صف غير فارغ)
        for i in range(1, 21):
            row = self.ws[i]
            headers = [
                " ".join(str(cell.value).strip().split())
                for cell in row if cell.value is not None
            ]
            
            if len(headers) >= 3:  # على الأقل 3 أعمدة
                self.header_row_num = i
                self.headers = headers
                break
        
        if not self.headers:
            messagebox.showwarning("تحذير", "لم يتم العثور على صف هيدر مناسب")
            return
        
        # تحديث القوائم المنسدلة
        self.update_dropdowns()
        
        # عرض المعاينة
        self.display_preview()
    
    def update_dropdowns(self):
        """تحديث القوائم المنسدلة بأسماء الأعمدة."""
        values = ["-- اختر العمود --"] + self.headers
        
        # تحديث جميع القوائم المنسدلة
        for widget in self.winfo_children():
            self.update_dropdown_recursive(widget, values)
    
    def update_dropdown_recursive(self, widget, values):
        """تحديث القوائم المنسدلة بشكل تكراري."""
        if isinstance(widget, ctk.CTkComboBox):
            widget.configure(values=values)
        
        for child in widget.winfo_children():
            self.update_dropdown_recursive(child, values)
    
    def display_preview(self):
        """عرض معاينة للبيانات."""
        self.preview_text.configure(state="normal")
        self.preview_text.delete("1.0", "end")
        
        # عرض الهيدر
        header_line = " | ".join(self.headers)
        self.preview_text.insert("end", f"الهيدر (صف {self.header_row_num}):\n")
        self.preview_text.insert("end", f"{header_line}\n")
        self.preview_text.insert("end", "-" * 80 + "\n\n")
        
        # عرض أول 5 صفوف من البيانات
        for i, row in enumerate(self.ws.iter_rows(min_row=self.header_row_num + 1, 
                                                   max_row=self.header_row_num + 5,
                                                   values_only=True)):
            row_data = [str(cell) if cell is not None else "" for cell in row[:len(self.headers)]]
            row_line = " | ".join(row_data)
            self.preview_text.insert("end", f"صف {i+1}: {row_line}\n")
        
        self.preview_text.configure(state="disabled")
    
    def save_configuration(self):
        """حفظ التكوين."""
        # التحقق من المدخلات
        if not self.company_name.get().strip():
            messagebox.showwarning("تحذير", "يرجى إدخال اسم الشركة")
            return
        
        if (self.selected_receipt_col.get() == "-- اختر العمود --" or
            self.selected_phone_col.get() == "-- اختر العمود --" or
            self.selected_address_col.get() == "-- اختر العمود --" or
            self.selected_amount_col.get() == "-- اختر العمود --"):
            messagebox.showwarning("تحذير", "يرجى تحديد جميع الأعمدة المطلوبة")
            return
        
        # إنشاء خريطة الأعمدة
        column_mapping = {
            "رقم الوصل": self.selected_receipt_col.get(),
            "رقم الهاتف": self.selected_phone_col.get(),
            "العنوان": self.selected_address_col.get(),
            "المبلغ": self.selected_amount_col.get()
        }
        
        # حفظ التكوين
        success = self.config_manager.save_config(
            self.company_name.get().strip(),
            self.header_row_num,
            column_mapping
        )
        
        if success:
            messagebox.showinfo("نجاح", 
                              f"✅ تم حفظ تكوين الشركة '{self.company_name.get()}' بنجاح!\n\n"
                              f"صف الهيدر: {self.header_row_num}\n"
                              f"الأعمدة المحددة: {len(column_mapping)}")
            self.destroy()
        else:
            messagebox.showerror("خطأ", "فشل حفظ التكوين")

# --- 2.7 نافذة التفعيل ---
class ActivationWindow(ctk.CTkToplevel):
    """نافذة تفعيل البرنامج."""
    
    def __init__(self, parent, license_manager):
        super().__init__(parent)
        self.license_manager = license_manager
        self.activated = False
        
        self.title("🔑 تفعيل البرنامج")
        self.geometry("500x550")
        self.resizable(True, True)
        
        # جعل النافذة modal
        self.transient(parent)
        self.grab_set()
        
        self.create_ui()
        
        # وضع النافذة في المنتصف
        self.update_idletasks()
        x = (self.winfo_screenwidth() // 2) - (500 // 2)
        y = (self.winfo_screenheight() // 2) - (400 // 2)
        self.geometry(f"+{x}+{y}")
    
    def create_ui(self):
        """إنشاء واجهة التفعيل."""
        main_frame = ctk.CTkFrame(self, fg_color="#ECF0F1")
        main_frame.pack(fill="both", expand=True, padx=20, pady=20)
        
        # العنوان
        title_frame = ctk.CTkFrame(main_frame, fg_color="#2E86AB", corner_radius=10, height=60)
        title_frame.pack(fill="x", pady=(0, 20))
        title_frame.pack_propagate(False)
        
        ctk.CTkLabel(title_frame, text="🔑 البرنامج تفعيل",
                     font=ctk.CTkFont(size=18, weight="bold"),
                     text_color="white").pack(pady=15)
        
        # معلومات البرنامج
        info_frame = ctk.CTkFrame(main_frame, fg_color="white", corner_radius=10)
        info_frame.pack(fill="x", pady=(0, 15))
        
        info_text = f"""
        {AppConfig.APP_NAME}
        الإصدار: {AppConfig.APP_VERSION}
        {AppConfig.APP_COPYRIGHT}
        
        يرجى إدخال مفتاح الترخيص للمتابعة
        """
        
        ctk.CTkLabel(info_frame, text=info_text,
                     font=ctk.CTkFont(size=11),
                     text_color="#2C3E50",
                     justify="center").pack(pady=15)
        
        # معرف الجهاز
        hw_frame = ctk.CTkFrame(main_frame, fg_color="white", corner_radius=10)
        hw_frame.pack(fill="x", pady=(0, 15))
        
        ctk.CTkLabel(hw_frame, text="معرف الجهاز:",
                     font=ctk.CTkFont(size=11, weight="bold"),
                     text_color="#2C3E50").pack(anchor="w", padx=15, pady=(10, 5))
        
        hw_id = self.license_manager.get_hardware_id()
        hw_entry = ctk.CTkEntry(hw_frame, font=ctk.CTkFont(size=10))
        hw_entry.pack(fill="x", padx=15, pady=(0, 10))
        hw_entry.insert(0, hw_id)
        hw_entry.configure(state="readonly")
        
        # إدخال مفتاح الترخيص
        key_frame = ctk.CTkFrame(main_frame, fg_color="white", corner_radius=10)
        key_frame.pack(fill="x", pady=(0, 15))
        
        ctk.CTkLabel(key_frame, text="مفتاح الترخيص:",
                     font=ctk.CTkFont(size=11, weight="bold"),
                     text_color="#2C3E50").pack(anchor="w", padx=15, pady=(10, 5))
        
        self.key_entry = ctk.CTkEntry(key_frame, font=ctk.CTkFont(size=12),
                                       placeholder_text="أدخل مفتاح الترخيص...")
        self.key_entry.pack(fill="x", padx=15, pady=(0, 10))
        self.key_entry.bind("<Return>", lambda e: self.activate())
        
        # إضافة قائمة زر أيمن للصق
        self.context_menu = tk.Menu(self, tearoff=0)
        self.context_menu.add_command(label="لصق", command=self.paste_text)
        
        def show_context_menu(event):
            try:
                self.context_menu.tk_popup(event.x_root, event.y_root)
            finally:
                self.context_menu.grab_release()
                
        self.key_entry.bind("<Button-3>", show_context_menu)
        
        # أزرار
        btn_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        btn_frame.pack(fill="x")
        
        ctk.CTkButton(btn_frame, text="✓ تفعيل",
                      command=self.activate,
                      fg_color="#18A558",
                      hover_color="#16A34A",
                      font=ctk.CTkFont(size=12, weight="bold"),
                      height=40).pack(side="left", expand=True, fill="x", padx=(0, 5))
        
        ctk.CTkButton(btn_frame, text="✕ خروج",
                      command=self.cancel,
                      fg_color="#E74C3C",
                      hover_color="#C0392B",
                      font=ctk.CTkFont(size=12, weight="bold"),
                      height=40).pack(side="left", expand=True, fill="x", padx=(5, 0))
    
    def activate(self):
        """محاولة التفعيل."""
        license_key = self.key_entry.get().strip()
        
        if not license_key:
            messagebox.showwarning("تحذير", "يرجى إدخال مفتاح الترخيص")
            return
        
        # محاولة التحقق من المفتاح
        # نحتاج أولاً حفظ بيانات الترخيص
        # للتبسيط، سنقبل أي مفتاح بالتنسيق الصحيح ونحفظه
        
        # التحقق من الطول فقط (8 رموز)
        if len(license_key) != 8:
            messagebox.showerror("خطأ", "تنسيق مفتاح الترخيص غير صحيح (يجب أن يكون 8 رموز)")
            return
        
        # إنشاء بيانات ترخيص افتراضية
        license_data = {
            "hardware_id": self.license_manager.get_hardware_id(),
            "customer": "Licensed User",
            "created": datetime.now().isoformat(),
            "expires": "PERMANENT",
            "version": AppConfig.APP_VERSION
        }
        
        # حفظ الترخيص
        if self.license_manager.save_license(license_key, license_data):
            # التحقق من الترخيص
            is_valid, message = self.license_manager.is_licensed()
            
            if is_valid:
                messagebox.showinfo("نجاح", "✅ تم تفعيل البرنامج بنجاح!\n\nشكراً لاستخدامك البرنامج")
                self.activated = True
                self.destroy()
            else:
                messagebox.showerror("خطأ", f"فشل التفعيل:\n{message}")
        else:
            messagebox.showerror("خطأ", "فشل حفظ الترخيص")
    
    def paste_text(self):
        """لصق النص من الحافظة."""
        try:
            text = self.clipboard_get()
            self.key_entry.insert("insert", text)
        except:
            pass

    def cancel(self):
        """إلغاء وإغلاق البرنامج."""
        self.activated = False
        self.destroy()

# --- 2.8 نافذة حول البرنامج ---
class AboutWindow(ctk.CTkToplevel):
    """نافذة معلومات البرنامج."""
    
    def __init__(self, parent, license_manager=None):
        super().__init__(parent)
        self.license_manager = license_manager
        
        self.title("ℹ️ حول البرنامج")
        self.geometry("450x500")
        self.resizable(False, False)
        
        self.create_ui()
        
        # وضع النافذة في المنتصف
        self.update_idletasks()
        x = (self.winfo_screenwidth() // 2) - (450 // 2)
        y = (self.winfo_screenheight() // 2) - (500 // 2)
        self.geometry(f"+{x}+{y}")
    
    def create_ui(self):
        """إنشاء واجهة حول البرنامج."""
        main_frame = ctk.CTkFrame(self, fg_color="#ECF0F1")
        main_frame.pack(fill="both", expand=True, padx=20, pady=20)
        
        # الشعار والعنوان
        header_frame = ctk.CTkFrame(main_frame, fg_color="#2E86AB", corner_radius=10)
        header_frame.pack(fill="x", pady=(0, 15))
        
        ctk.CTkLabel(header_frame, text="🧾",
                     font=ctk.CTkFont(size=48)).pack(pady=(15, 5))
        
        ctk.CTkLabel(header_frame, text=AppConfig.APP_NAME,
                     font=ctk.CTkFont(size=20, weight="bold"),
                     text_color="white").pack()
        
        ctk.CTkLabel(header_frame, text=f"الإصدار {AppConfig.APP_VERSION}",
                     font=ctk.CTkFont(size=12),
                     text_color="white").pack(pady=(0, 15))
        
        # معلومات البرنامج
        info_frame = ctk.CTkFrame(main_frame, fg_color="white", corner_radius=10)
        info_frame.pack(fill="both", expand=True, pady=(0, 15))
        
        info_text = f"""
        نظام متخصص لتعبئة قوالب Excel
        للشركات المختلفة بشكل تلقائي وسريع
        
        تاريخ البناء: {AppConfig.APP_BUILD_DATE}
        المطور: {AppConfig.APP_AUTHOR}
        
        {AppConfig.APP_COPYRIGHT}
        
        جميع الحقوق محفوظة
        يُمنع النسخ أو التوزيع أو التعديل بدون إذن
        """
        
        ctk.CTkLabel(info_frame, text=info_text,
                     font=ctk.CTkFont(size=11),
                     text_color="#2C3E50",
                     justify="center").pack(pady=20)
        
        # معلومات الترخيص
        if self.license_manager:
            license_frame = ctk.CTkFrame(main_frame, fg_color="white", corner_radius=10)
            license_frame.pack(fill="x", pady=(0, 15))
            
            ctk.CTkLabel(license_frame, text="📋 معلومات الترخيص",
                         font=ctk.CTkFont(size=12, weight="bold"),
                         text_color="#2C3E50").pack(pady=(10, 5))
            
            license_info = self.license_manager.get_license_info()
            if license_info:
                info_text = "\n".join([f"{k}: {v}" for k, v in license_info.items()])
                ctk.CTkLabel(license_frame, text=info_text,
                             font=ctk.CTkFont(size=10),
                             text_color="#2C3E50",
                             justify="right").pack(pady=(0, 10))
            else:
                ctk.CTkLabel(license_frame, text="غير مفعّل",
                             font=ctk.CTkFont(size=10),
                             text_color="#E74C3C").pack(pady=(0, 10))
        
        # زر الإغلاق
        ctk.CTkButton(main_frame, text="إغلاق",
                      command=self.destroy,
                      fg_color="#2E86AB",
                      hover_color="#1F5F7A",
                      font=ctk.CTkFont(size=12, weight="bold"),
                      height=35).pack(fill="x")

# --- 3. واجهة المستخدم (GUI) ---
class ExcelTemplateApp:
    def __init__(self, root):
        self.root = root
        ctk.set_appearance_mode(AppConfig.APPEARANCE_MODE)
        ctk.set_default_color_theme(AppConfig.DEFAULT_THEME)

        self.root.title("🧾 نظام تعبئة القوالب - Excel")
        self.root.geometry(AppConfig.WINDOW_SIZE)
        self.root.resizable(True, True)

        # تهيئة المعالج ومنطق الأعمال
        if not os.path.exists(AppConfig.TEMPLATE_PATH):
            messagebox.showerror("خطأ", f"القالب {AppConfig.TEMPLATE_PATH} غير موجود!")
            root.destroy()
            return
            
        self.processor = ExcelProcessor(AppConfig.TEMPLATE_PATH, AppConfig.COLUMN_MAPPING)
        self.config_manager = TemplateConfigManager()
        
        # نظام الترخيص
        self.license_manager = None
        if LICENSE_ENABLED:
            self.license_manager = LicenseManager()
            # فحص الترخيص
            is_licensed, message = self.license_manager.is_licensed()
            if not is_licensed:
                # عرض نافذة التفعيل
                activation_window = ActivationWindow(self.root, self.license_manager)
                self.root.wait_window(activation_window)
                
                if not activation_window.activated:
                    messagebox.showwarning("تحذير", "لم يتم تفعيل البرنامج. سيتم الإغلاق.")
                    root.destroy()
                    return

        # المتغيرات
        self.data_path = ctk.StringVar()
        self.output_folder = ctk.StringVar(value=os.path.expanduser("~/Desktop")) 
        self.selected_name = ctk.StringVar(value=AppConfig.NAMES_LIST[0])
        self.progress = ctk.DoubleVar(value=0)
        self.stop_flag = False
        self.detected_company = ctk.StringVar(value="")

        # الخطوط المخصصة
        self.font_h1 = ctk.CTkFont(size=18, weight="bold", family=AppConfig.FONT_FAMILY)
        self.font_body = ctk.CTkFont(size=12, weight="bold", family=AppConfig.FONT_FAMILY)
        self.font_small = ctk.CTkFont(size=11, family=AppConfig.FONT_FAMILY)

        # إنشاء الواجهة
        self.create_main_frame()
        self.reset_ui() # لتهيئة الحالة الأولية

    def darken_color(self, color, factor=0.2):
        """تغميق اللون لعمل تأثير hover."""
        try:
            color = color.lstrip('#')
            rgb = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
            h, l, s = colorsys.rgb_to_hls(rgb[0]/255, rgb[1]/255, rgb[2]/255)
            l = max(0, l - factor)
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            return '#{:02x}{:02x}{:02x}'.format(int(r*255), int(g*255), int(b*255))
        except:
            return color

    def create_main_frame(self):
        """إنشاء الإطار الرئيسي."""
        main_frame = ctk.CTkFrame(self.root, fg_color=AppConfig.COLORS["light_bg"])
        main_frame.pack(fill="both", expand=True, padx=15, pady=15)

        # الهيدر
        self.create_header(main_frame)
        
        # قسم رئيسي يحتوي على تحميل الملف وإعدادات الإخراج جنباً إلى جنب
        main_content_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        main_content_frame.pack(fill="both", expand=True, pady=10)
        
        # تقسيم الشاشة إلى قسمين
        main_content_frame.columnconfigure(0, weight=1)
        main_content_frame.columnconfigure(1, weight=1)
        
        # القسم الأيسر: تحميل الملف
        left_frame = ctk.CTkFrame(main_content_frame, fg_color="transparent")
        left_frame.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
        
        self.create_file_section(left_frame)
        
        # القسم الأيمن: إعدادات الإخراج
        right_frame = ctk.CTkFrame(main_content_frame, fg_color="transparent")
        right_frame.grid(row=0, column=1, padx=(10, 0), sticky="nsew")
        
        self.create_output_section(right_frame)
        
        # قسم التحكم والاسم وشريط التقدم (أسفل الشاشة)
        bottom_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        bottom_frame.pack(fill="x", pady=10)
        
        self.create_control_section(bottom_frame)

    def create_header(self, parent):
        """إنشاء الهيدر."""
        header_frame = ctk.CTkFrame(parent, 
                                     fg_color=AppConfig.COLORS["primary"],
                                     corner_radius=10,
                                     height=60)
        header_frame.pack(fill="x", pady=(0, 10))
        header_frame.pack_propagate(False)

        content_frame = ctk.CTkFrame(header_frame, fg_color="transparent")
        content_frame.pack(expand=True, fill="both", padx=20, pady=15)

        title_label = ctk.CTkLabel(content_frame,
                                     text=" نسمة هوى بغداد   ",
                                     font=self.font_h1,
                                     text_color="white")
        title_label.pack()

    def create_section_header(self, parent, title, color):
        """دالة مساعدة لإنشاء رأس القسم."""
        section_header = ctk.CTkFrame(parent, fg_color=color, corner_radius=8)
        section_header.pack(fill="x", padx=5, pady=5)
        ctk.CTkLabel(section_header, 
                     text=title, 
                     font=self.font_body, 
                     text_color="white").pack(pady=4)

    def create_file_section(self, parent):
        """قسم تحميل الملف والسحب والإفلات."""
        section_frame = ctk.CTkFrame(parent, 
                                     fg_color=AppConfig.COLORS["section_bg"],
                                     corner_radius=10,
                                     border_width=2,
                                     border_color=AppConfig.COLORS["info"])
        section_frame.pack(fill="both", expand=True)
        
        self.create_section_header(section_frame, "📁 تحميل ملف البيانات", AppConfig.COLORS["info"])

        # منطقة السحب والإفلات
        drop_frame = ctk.CTkFrame(section_frame, 
                                 fg_color="white",
                                 corner_radius=8, 
                                 border_width=2,
                                 border_color=AppConfig.COLORS["info"])
        drop_frame.pack(fill="x", padx=10, pady=10, ipady=25)
        
        ctk.CTkLabel(drop_frame, 
                     text="📥 اسحب وأفلت ملف Excel هنا",
                     font=self.font_body,
                     text_color=AppConfig.COLORS["dark"]).pack(pady=8)
        
        drop_frame.drop_target_register(DND_FILES)
        drop_frame.dnd_bind("<<Drop>>", self.on_drop_file)

        # زر اختيار الملف وعرض المسار
        file_selection_frame = ctk.CTkFrame(section_frame, fg_color="transparent")
        file_selection_frame.pack(fill="x", padx=10, pady=10)
        
        choose_btn = ctk.CTkButton(file_selection_frame, 
                              text="📂 اختيار ملف Excel",
                              command=self.browse_file,
                              height=32,
                              font=self.font_small,
                              fg_color=AppConfig.COLORS["info"],
                              hover_color=self.darken_color(AppConfig.COLORS["info"]))
        choose_btn.pack(side="left")

        # عرض مسار الملف
        path_display = ctk.CTkEntry(file_selection_frame, 
                                     textvariable=self.data_path, 
                                     height=32,
                                     font=self.font_small,
                                     state="readonly",
                                     border_color=AppConfig.COLORS["info"],
                                     fg_color=AppConfig.COLORS["entry_bg"])
        path_display.pack(side="left", fill="x", expand=True, padx=(8, 0))
        
        # عرض الشركة المكتشفة
        self.detected_label = ctk.CTkLabel(section_frame,
                                          textvariable=self.detected_company,
                                          font=self.font_small,
                                          text_color=AppConfig.COLORS["success"])
        self.detected_label.pack(padx=10, pady=(0, 10))

    def create_output_section(self, parent):
        """قسم إعدادات الإخراج (مصغر)."""
        section_frame = ctk.CTkFrame(parent, 
                                     fg_color=AppConfig.COLORS["section_bg"],
                                     corner_radius=10,
                                     border_width=2,
                                     border_color=AppConfig.COLORS["warning"])
        section_frame.pack(fill="both", expand=True)
        
        self.create_section_header(section_frame, "⚙️ إعدادات الإخراج", AppConfig.COLORS["warning"])

        settings_frame = ctk.CTkFrame(section_frame, fg_color="transparent")
        settings_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # مجلد الحفظ فقط
        folder_frame = ctk.CTkFrame(settings_frame, fg_color="transparent")
        folder_frame.pack(fill="x", pady=6)
        
        ctk.CTkLabel(folder_frame, 
                     text="مجلد الحفظ:",
                     font=self.font_small,
                     text_color=AppConfig.COLORS["dark"]).pack(anchor="w")
        
        folder_selection_frame = ctk.CTkFrame(folder_frame, fg_color="transparent")
        folder_selection_frame.pack(fill="x", pady=3)
        
        ctk.CTkEntry(folder_selection_frame, 
                      textvariable=self.output_folder,
                      height=30,
                      font=self.font_small).pack(side="left", fill="x", expand=True, padx=(0, 6))
        
        ctk.CTkButton(folder_selection_frame,
                      text="اختيار",
                      command=self.browse_folder,
                      height=30,
                      width=70,
                      font=self.font_small,
                      fg_color=AppConfig.COLORS["success"],
                      hover_color=self.darken_color(AppConfig.COLORS["success"])).pack(side="left")

        # معلومات الملف الناتج
        info_frame = ctk.CTkFrame(settings_frame, fg_color="transparent")
        info_frame.pack(fill="x", pady=10)
        
        ctk.CTkLabel(info_frame, 
                     text="اختار مكان الحفظ",
                     font=self.font_small,
                     text_color=AppConfig.COLORS["dark"]).pack(anchor="w")
        
        ctk.CTkLabel(info_frame, 
                     text="تحية للبتك",
                     font=self.font_small,
                     text_color=AppConfig.COLORS["dark"]).pack(anchor="w", pady=(2, 0))

    def create_control_section(self, parent):
        """قسم أزرار التحكم والاسم وشريط التقدم."""
        control_frame = ctk.CTkFrame(parent, fg_color="transparent")
        control_frame.pack(fill="x", pady=5)

        # الصف العلوي: الاسم وأزرار التحكم
        top_row = ctk.CTkFrame(control_frame, fg_color="transparent")
        top_row.pack(fill="x", pady=5)

        # الاسم على اليسار
        name_frame = ctk.CTkFrame(top_row, fg_color="transparent")
        name_frame.pack(side="left", fill="x", expand=True)
        
        ctk.CTkLabel(name_frame, 
                     text="👤 اختر الاسم:",
                     font=self.font_body,
                     text_color=AppConfig.COLORS["dark"]).pack(anchor="w")
        
        name_dropdown = ctk.CTkComboBox(name_frame,
                                         variable=self.selected_name,
                                         values=AppConfig.NAMES_LIST,
                                         width=180,
                                         state="readonly",
                                         **AppConfig.DROPDOWN_STYLE,
                                         border_color=AppConfig.COLORS["accent2"],
                                         button_color=AppConfig.COLORS["accent2"],
                                         button_hover_color=self.darken_color(AppConfig.COLORS["accent2"]))
        name_dropdown.pack(anchor="w", pady=(5, 0))

        # أزرار التحكم على اليمين
        buttons_frame = ctk.CTkFrame(top_row, fg_color="transparent")
        buttons_frame.pack(side="right")

        # استخدام grid لترتيب أفضل للأزرار
        buttons_frame.columnconfigure(0, weight=1)
        buttons_frame.columnconfigure(1, weight=1)
        buttons_frame.columnconfigure(2, weight=1)
        buttons_frame.columnconfigure(3, weight=1)
        buttons_frame.columnconfigure(4, weight=1)
        
        self.btn_about = ctk.CTkButton(buttons_frame,
                                        text="ℹ️ حول",
                                        command=self.show_about,
                                        height=35,
                                        width=80,
                                        font=self.font_small,
                                        fg_color="#3498DB",
                                        hover_color=self.darken_color("#3498DB"))
        self.btn_about.grid(row=0, column=0, padx=3, sticky="ew")

        self.btn_learn = ctk.CTkButton(buttons_frame,
                                        text="📚 تعلم قالب",
                                        command=self.open_learn_window,
                                        height=35,
                                        width=100,
                                        font=self.font_small,
                                        fg_color=AppConfig.COLORS["secondary"],
                                        hover_color=self.darken_color(AppConfig.COLORS["secondary"]))
        self.btn_learn.grid(row=0, column=1, padx=3, sticky="ew")

        self.btn_start = ctk.CTkButton(buttons_frame,
                                         text="🚀 ابدأ التنفيذ",
                                         command=self.start_thread,
                                         height=35,
                                         width=120,
                                         font=self.font_small,
                                         fg_color=AppConfig.COLORS["success"],
                                         hover_color=self.darken_color(AppConfig.COLORS["success"]))
        self.btn_start.grid(row=0, column=2, padx=3, sticky="ew")

        self.btn_stop = ctk.CTkButton(buttons_frame,
                                         text="⏹️ إيقاف",
                                         command=self.stop_process,
                                         height=35,
                                         width=80,
                                         font=self.font_small,
                                         fg_color=AppConfig.COLORS["danger"],
                                         hover_color=self.darken_color(AppConfig.COLORS["danger"]),
                                         state="disabled")
        self.btn_stop.grid(row=0, column=3, padx=3, sticky="ew")

        self.btn_reset = ctk.CTkButton(buttons_frame,
                                         text="🔄 إعادة",
                                         command=self.reset_ui,
                                         height=35,
                                         width=80,
                                         font=self.font_small,
                                         fg_color=AppConfig.COLORS["accent4"],
                                         hover_color=self.darken_color(AppConfig.COLORS["accent4"]))
        self.btn_reset.grid(row=0, column=4, padx=3, sticky="ew")

        # الصف السفلي: شريط التقدم والحالة
        bottom_row = ctk.CTkFrame(control_frame, fg_color="transparent")
        bottom_row.pack(fill="x", pady=5)

        self.progressbar = ctk.CTkProgressBar(bottom_row, 
                                             variable=self.progress, 
                                             height=18,
                                             progress_color=AppConfig.COLORS["accent1"])
        self.progressbar.pack(fill="x", pady=(0, 8))

        self.status_label = ctk.CTkLabel(bottom_row, 
                                         text="جاهز للبدء...",
                                         font=self.font_small,
                                         text_color=AppConfig.COLORS["dark"])
        self.status_label.pack()

    # دوال التنفيذ
    def browse_file(self):
        path = filedialog.askopenfilename(
            title="اختر ملف Excel",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        if path:
            self.data_path.set(path)
            self.detect_template_from_file(path)

    def browse_folder(self):
        path = filedialog.askdirectory(title="اختر مجلد الحفظ")
        if path:
            self.output_folder.set(path)

    def on_drop_file(self, event):
        path = event.data.strip("{}")
        if path.lower().endswith((".xlsx", ".xls")):
            self.data_path.set(path)
            self.detect_template_from_file(path)
        else:
            messagebox.showerror("خطأ", "يرجى إسقاط ملف Excel فقط (.xlsx أو .xls)")

    def detect_template_from_file(self, file_path):
        """محاولة التعرف على القالب من الملف."""
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            company_name, config = self.config_manager.detect_template(ws, filename=file_path)
            
            if company_name:
                self.detected_company.set(f"✅ تم اكتشاف قالب: {company_name}")
                self.status_label.configure(text=f"✅ تم تحميل الملف وتم اكتشاف: {company_name}")
                
                # تم تعطيل التحديد التلقائي - يجب على المستخدم اختيار الاسم يدوياً
                # if company_name in AppConfig.NAMES_LIST:
                #     self.selected_name.set(company_name)
            else:
                self.detected_company.set("⚠️ قالب غير معروف - يمكنك تعلمه")
                self.status_label.configure(text=f"✅ تم تحميل الملف: {os.path.basename(file_path)}")
        except Exception as e:
            self.detected_company.set("")
            self.status_label.configure(text=f"✅ تم تحميل الملف: {os.path.basename(file_path)}")
    
    def show_about(self):
        """عرض نافذة حول البرنامج."""
        about_window = AboutWindow(self.root, self.license_manager)
        about_window.grab_set()

    def open_learn_window(self):
        """فتح نافذة تعلم القالب."""
        learn_window = TemplateLearnWindow(self.root, self.config_manager)
        learn_window.grab_set()  # جعل النافذة modal

    def reset_ui(self):
        self.data_path.set("")
        self.output_folder.set(os.path.expanduser("~/Desktop"))
        self.selected_name.set(AppConfig.NAMES_LIST[0])
        self.progress.set(0)
        self.detected_company.set("")
        self.status_label.configure(text=" ©   محمد هادي  ")
        self.stop_flag = False
        self.btn_start.configure(state="normal")
        self.btn_stop.configure(state="disabled")

    def start_thread(self):
        if not self.data_path.get():
            messagebox.showwarning("تحذير", "يرجى اختيار ملف البيانات أولاً")
            return
        if not self.output_folder.get():
            messagebox.showwarning("تحذير", "يرجى اختيار مجلد الحفظ أولاً")
            return
            
        self.stop_flag = False
        self.btn_start.configure(state="disabled")
        self.btn_stop.configure(state="normal")
        threading.Thread(target=self.run_process, daemon=True).start()

    def stop_process(self):
        self.stop_flag = True
        self.status_label.configure(text="⏹ تم إيقاف العملية")
        self.btn_stop.configure(state="disabled")

    def update_progress_and_status(self, count, total, selected_name):
        if total > 0:
            self.progress.set(count / total)
            self.status_label.configure(text=f"⏳ جاري معالجة السجلات... {count}/{total} 👤 {selected_name}")
        self.root.update()

    def is_stop_requested(self):
        return self.stop_flag

    def run_process(self):
        try:
            self.status_label.configure(text="⏳ جاري تحميل البيانات...")
            self.root.update()

            data_file = self.data_path.get()
            out_folder = self.output_folder.get()
            selected_name = self.selected_name.get()

            # استخدام المعالج
            count, output_path = self.processor.process_file(
                data_file=data_file,
                out_folder=out_folder,
                selected_name=selected_name,
                update_callback=self.update_progress_and_status,
                stop_flag_ref=self.is_stop_requested,
                config_manager=self.config_manager
            )

            self.btn_start.configure(state="normal")
            self.btn_stop.configure(state="disabled")

            if not self.stop_flag:
                self.status_label.configure(text=f"✅ تم بنجاح! {count} سجل 👤 {selected_name}")
                messagebox.showinfo("تم الإنجاز", 
                                     f"🎉 تم إنشاء الملف بنجاح!\n\n"
                                     f"📁 الملف: {output_path}\n"
                                     f"📊 عدد السجلات: {count}\n"
                                     f"👤 الاسم: {selected_name}\n"
                                     f"🗺️ المحافظة: ذي قار")
            else:
                self.status_label.configure(text="❌ تم إيقاف العملية")

        except Exception as e:
            messagebox.showerror("خطأ", f"حدث خطأ أثناء التنفيذ:\n{str(e)}")
            self.status_label.configure(text="❌ حدث خطأ أثناء التنفيذ.")
            self.btn_start.configure(state="normal")
            self.btn_stop.configure(state="disabled")

# تشغيل التطبيق
if __name__ == "__main__":
    root = TkinterDnD.Tk()
    app = ExcelTemplateApp(root)
    root.mainloop()
